from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

FONT = "MiSans"
INK = RGBColor(0x1D, 0x1D, 0x1A)
BODY = RGBColor(0x4D, 0x4D, 0x4D)
MUTE = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0xA3, 0xA3, 0xA3)
RED = RGBColor(0xC0, 0x00, 0x00)
NAVY = RGBColor(0x00, 0x20, 0x60)
PALE = RGBColor(0xF2, 0xF2, 0xF2)
RULE = RGBColor(0xDD, 0xDD, 0xDD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_SOFT = RGBColor(0xA9, 0xB8, 0xD8)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_run_font(run, size, bold, color, italic=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def tb(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=3, line=1.12):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space)
        p.line_spacing = line
        runs = para if isinstance(para, list) else [para]
        for spec in runs:
            text, size, bold, color = spec[0], spec[1], spec[2], spec[3]
            italic = spec[4] if len(spec) > 4 else False
            r = p.add_run()
            r.text = text
            set_run_font(r, size, bold, color, italic)
    return box


def rect(slide, x, y, w, h, fill, line_color=None, line_w=1.0, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    return sh


def hrule(slide, x, y, w, color=RULE):
    rect(slide, x, y, w, 0.012, color)


def sec_head(slide, x, y, w, text, red=False, size=13):
    tb(slide, x, y, w, 0.26, [(text, size, True, RED if red else INK)])
    hrule(slide, x, y + 0.3, w)


def header(slide, title, tag, subtitle=None):
    tb(slide, 0.67, 0.36, 9.7, 0.5, [(title, 25, True, INK)])
    tb(slide, 9.4, 0.5, 3.27, 0.24, [(tag, 10, False, LIGHT)], align=PP_ALIGN.RIGHT)
    rect(slide, 0.67, 0.94, 0.64, 0.04, RED)
    if subtitle:
        tb(slide, 0.67, 1.06, 12.0, 0.26, [(subtitle, 11.5, False, MUTE)])


def footer(slide, n):
    tb(slide, 11.75, 7.2, 0.92, 0.22, [(n + " / 8", 9, False, LIGHT)], align=PP_ALIGN.RIGHT)


def navy_bar(slide, y, paras, h=0.42):
    rect(slide, 0.67, y, 12.0, h, NAVY)
    tb(slide, 0.97, y, 11.4, h, paras, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=0)


def make_table(slide, x, y, w, col_widths, row_heights):
    g = slide.shapes.add_table(len(row_heights), len(col_widths), Inches(x), Inches(y), Inches(w), Inches(sum(row_heights)))
    t = g.table
    t.first_row = False
    t.horz_banding = False
    tblPr = t._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    for i, cw in enumerate(col_widths):
        t.columns[i].width = Inches(cw)
    for i, rh in enumerate(row_heights):
        t.rows[i].height = Inches(rh)
    return t


def set_cell(t, r, c, segs, size, fill, align=PP_ALIGN.LEFT, line=1.05):
    cell = t.cell(r, c)
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.1)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line
    if not isinstance(segs, list):
        segs = [(segs, False, BODY)]
    for spec in segs:
        text, bold, color = spec[0], spec[1], spec[2]
        r_ = p.add_run()
        r_.text = text
        set_run_font(r_, size, bold, color)


def arrow(slide, x, y, w=0.26, h=0.3, size=13):
    tb(slide, x, y, w, h, [("\u2192", size, False, LIGHT)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def cover():
    s = prs.slides.add_slide(BLANK)
    rect(s, 8.33, 0, 5.0, 7.5, NAVY)
    rect(s, 0.67, 1.78, 0.64, 0.04, RED)
    tb(s, 0.67, 1.98, 7.2, 0.3, [("研究课题立项汇报 · 2026.09", 12, True, RED)])
    tb(s, 0.67, 2.47, 7.22, 1.39, [
        ("面向智能体工作负载的", 32, True, INK),
        ("轨迹级查询处理与优化", 32, True, INK),
    ], space=2)
    tb(s, 0.67, 4.06, 7.22, 0.61, [
        ("Agent-Aware Trajectory-Level Query Processing", 15, False, MUTE),
        ("and Optimization", 15, False, MUTE),
    ], space=1)
    tb(s, 0.67, 4.94, 7.22, 0.36, [("从单条 SQL 优化,走向持续演化的 Agent 查询轨迹优化。", 13, False, BODY)])
    tb(s, 0.67, 5.42, 7.22, 0.3, [("Future Query Unknown, Future Intent Visible.", 14, True, RED)])
    tb(s, 0.67, 6.94, 6.94, 0.25, [("汇报人:___　|　部门:___　|　2026-09", 10, False, LIGHT)])

    tb(s, 8.95, 2.25, 4.2, 0.4, [("\u03c4 = \u27e8q\u2081, r\u2081, \u2026, q\u1d40, r\u1d40\u27e9", 20, True, WHITE)])
    tb(s, 8.95, 2.72, 4.2, 0.28, [("同一条结果驱动的查询轨迹", 11.5, False, NAVY_SOFT)])
    rect(s, 8.95, 3.35, 0.8, 0.016, RGBColor(0x3A, 0x5A, 0x9E))
    tb(s, 8.95, 3.72, 4.2, 0.45, [("Query Unknown,", 24, True, WHITE)])
    tb(s, 8.95, 4.28, 4.2, 0.45, [("Intent Visible.", 24, True, WHITE)])
    tb(s, 8.95, 5.1, 4.4, 0.28, [("优化范式:Reactive \u2192 Predictive \u2192 Cooperative", 11, False, NAVY_SOFT)])


def page1():
    s = prs.slides.add_slide(BLANK)
    header(s, "Agent 正在改变数据库的使用方式", "01 · 背景趋势",
           "海内外数据库产品都在把「用数据的人」换成 Agent \u2014\u2014 但数据库内核还没有为此改变")

    sec_head(s, 0.67, 1.42, 5.83, "海外 · 2025")
    overseas = [
        ("Snowflake · Cortex Agents(2025.11 GA)", "全托管 Agent 平台:NL\u2192SQL / 检索 / 代码执行,规划\u2014工具\u2014反思循环"),
        ("Databricks · Genie + Deep Research", "对话式 BI;\u201c为什么\u201d类问题自动规划研究、并行验证假设并归因"),
        ("Google BigQuery · Data Agents", "对话式分析;verified queries + 参数化 SQL 模板保证业务口径"),
        ("Microsoft Fabric · Copilot(Agent Mode)", "自然语言生成 T-SQL、执行计划分析与优化建议、多步工具调用"),
    ]
    yy = 1.82
    for name, desc in overseas:
        tb(s, 0.67, yy, 5.83, 0.22, [(name, 11, True, INK)])
        tb(s, 0.67, yy + 0.215, 5.83, 0.2, [(desc, 9.5, False, BODY)], space=0)
        yy += 0.47

    sec_head(s, 6.84, 1.42, 5.83, "国内 · 2025\u20132026")
    domestic = [
        ("阿里云 · Data Agent for Analytics / ADBAgent", "企业级数据分析智能体:意图理解\u2192规划\u2192分析\u2192报告,沉淀行业 Skills"),
        ("OceanBase · AI 数据库(Lakebase)", "Agent 记忆、上下文工程、数据沙箱与库级克隆,支撑多 Agent 并行试错"),
        ("TiDB · AI Agent Database", "数据库 Branching、持久记忆、百万级 schema:\u201c为每个 Agent 配一个库\u201d"),
        ("腾讯云 · DatabaseClaw / TDSQL-C AI 应用", "自然语言运维:MySQL 工单覆盖 28.4%,CPU 异常诊断 30min\u21922min"),
        ("华为 · GaussDB(AI-Native)", "智能问答 / NL2SQL / 智能诊断,率先通过信通院 AI 数据库工具测评"),
    ]
    yy = 1.82
    for name, desc in domestic:
        tb(s, 6.84, yy, 5.83, 0.22, [(name, 10.5, True, INK)])
        tb(s, 6.84, yy + 0.215, 5.83, 0.2, [(desc, 9.5, False, BODY)], space=0)
        yy += 0.47

    t = make_table(s, 0.67, 4.24, 12.0, [1.75, 5.85, 4.4], [0.34] + [0.35] * 5)
    set_cell(t, 0, 0, [("布局层面", True, RED)], 11.5, None)
    set_cell(t, 0, 1, [("厂商在做什么,图什么", True, RED)], 11.5, None)
    set_cell(t, 0, 2, [("代表产品", True, RED)], 11.5, None)
    rows = [
        ("接口层 · NL2SQL / ChatBI", "降低用数门槛:用户面从「会 SQL 的人」扩大到全部业务人员",
         "BigQuery Data Agents · Fabric Copilot · ADBAgent"),
        ("编排层 · 多步推理 / 根因归因", "把分析师的探索过程自动化:从被动问答升级为自主交付洞察",
         "Cortex Agents · Genie + Deep Research · Data Agent for Analytics"),
        ("运维层 · NL 运维 / 智能诊断", "运维自治、降低长期总成本(TCO),缓解 DBA 短缺",
         "DatabaseClaw / TDSQL-C AI · GaussDB"),
        ("平台层 · 记忆 / 沙箱 / Branching", "成为 Agent 应用的运行时底座:Git 式秒级分支,供并行试错与回滚",
         "OceanBase Lakebase · TiDB"),
        ("内核层 · \u2014\u2014 空白 \u2014\u2014", "查询处理仍假设 SQL 逐条独立到达、独立优化,对轨迹无感知",
         "\u2190 本课题:轨迹级查询处理与优化"),
    ]
    for i, (a, b, c) in enumerate(rows):
        last = i == 4
        fill = PALE if last else (PALE if i % 2 == 1 else WHITE)
        col = RED if last else BODY
        set_cell(t, i + 1, 0, [(a, True, RED if last else INK)], 10.5, fill)
        set_cell(t, i + 1, 1, [(b, last, col)], 10, fill)
        set_cell(t, i + 1, 2, [(c, True, col)], 10, fill)

    navy_bar(s, 6.72, [[
        ("外围四层都在拥抱 Agent,唯独查询处理内核仍是 Agent-Blind \u2014\u2014 ", 12, True, WHITE),
        ("本课题要补上内核这一层", 12, True, RGBColor(0xFF, 0xB3, 0xB3)),
    ]])
    tb(s, 0.67, 7.2, 9.5, 0.2,
       [("来源:Snowflake / Databricks / Google Cloud / Microsoft Learn 官方文档;阿里云、OceanBase、PingCAP、腾讯云、华为官方资料(2025\u20132026)", 8.5, False, LIGHT)])
    footer(s, "2")


def page2():
    s = prs.slides.add_slide(BLANK)
    header(s, "Agent 的数据访问不是一条 SQL,而是一条结果驱动的查询轨迹", "02 · 问题背景",
           "一个高层任务 \u2192 一串相互依赖的 SQL:对数据库而言是一种全新的 workload 形态")

    boxes = [
        (1.60, 1.20, "任务", "为什么 8 月\n收入下降?", PALE, None, INK, True),
        (3.06, 1.45, "q\u2081", "按国家统计\n8 月收入", WHITE, RULE, INK, True),
        (4.77, 1.20, "r\u2081", "加拿大\n收入下降", WHITE, RULE, MUTE, True),
        (6.23, 1.55, "Agent", "依据结果推理\n决定下一条", RED, None, WHITE, True),
        (8.04, 1.45, "q\u2082", "加拿大 \u00d7\n产品分析", WHITE, RULE, INK, True),
        (9.75, 0.40, "", "\u22ef", None, None, MUTE, True),
        (10.41, 1.30, "q\u1d40", "时间趋势\n验证", WHITE, RULE, INK, True),
    ]
    for x, w, tag, body, fill, ln, tcol, bold in boxes:
        if fill is not None or ln is not None:
            rect(s, x, 1.38, w, 0.85, fill, ln, 1.0)
        if tag:
            tb(s, x, 1.44, w, 0.2, [(tag, 9, True, WHITE if fill == RED else RED)], align=PP_ALIGN.CENTER)
        tb(s, x, 1.66, w, 0.52, [(body, 9.5, bold, tcol)], align=PP_ALIGN.CENTER, space=0, line=1.1)
    for ax in (2.80, 4.51, 5.97, 7.78, 9.49, 10.15):
        arrow(s, ax, 1.66)

    tb(s, 0.67, 2.44, 12.0, 0.24, [[
        ("q\u209c\u208a\u2081 = Agent(Task, State\u209c, q\u208a\u209c, r\u208a\u209c)", 11, True, INK),
        ("   \u2014\u2014 下一条 SQL 由前序结果与任务状态共同决定;r\u2081 若是别的结果,后续轨迹随之完全不同", 11, False, MUTE),
    ]], align=PP_ALIGN.CENTER)

    rows = [
        ("对比维度", "人 / 传统应用程序", "Data Agent(智能体)"),
        ("访问模式", "手写 SQL 或预定义模板,一次提交", "围绕任务连续生成多步 SQL,迭代式访问"),
        ("查询间关系", "彼此独立或弱相关", "数据范围、谓词、Join、聚合维度高度相关"),
        ("下一查询由谁决定", "人 / 固定业务逻辑", "由前序结果动态决定"),
        ("未来意图", "不可见:存在于人脑或黑盒程序中", "可主动暴露:Goal / Plan / Stage / Candidate Actions"),
        ("访问空间", "受 UI、报表模板限制", "任意 SQL、跨表 Join、schema 探索,开放且难预测"),
        ("合适的优化单元", "单条查询", "完整查询轨迹 \u03c4"),
    ]
    t = make_table(s, 0.67, 2.86, 12.0, [2.3, 4.6, 5.1], [0.36] + [0.4] * 6)
    for i, (a, b, c) in enumerate(rows):
        if i == 0:
            set_cell(t, 0, 0, [(a, True, RED)], 11.5, None)
            set_cell(t, 0, 1, [(b, True, RED)], 11.5, None, PP_ALIGN.CENTER)
            set_cell(t, 0, 2, [(c, True, RED)], 11.5, None, PP_ALIGN.CENTER)
        else:
            fill = PALE if i % 2 == 0 else WHITE
            set_cell(t, i, 0, [(a, True, INK)], 10.5, fill)
            set_cell(t, i, 1, [(b, False, BODY)], 10, fill)
            set_cell(t, i, 2, [(c, False, INK)], 10, fill)

    navy_bar(s, 6.72, [[
        ("\u672a\u6765\u67e5\u8be2\u672a\u77e5\uff0c\u4f46\u672a\u6765\u610f\u56fe\u90e8\u5206\u53ef\u89c1 \u2014\u2014 ", 12, True, WHITE),
        ("Query Unknown, Intent Visible", 12, True, RGBColor(0xFF, 0xB3, 0xB3)),
        (" \u2014\u2014 \u4e00\u79cd\u6570\u636e\u5e93\u4ece\u672a\u9762\u5bf9\u8fc7\u7684\u4fe1\u606f\u6761\u4ef6", 12, True, WHITE),
    ]])
    footer(s, "3")


def page3():
    s = prs.slides.add_slide(BLANK)
    header(s, "单查询最优 \u2260 Agent 任务最优:一个带数字的例子", "02 · 问题背景",
           "为什么逐条独立优化会输 \u2014\u2014 sales 表 20 GB / 1 亿行,8 月子集 10%(2 GB),代价单位任意")

    tb(s, 0.67, 1.44, 12.0, 0.24, [[
        ("q\u2081", 11, True, RED), (":8 月 \u00d7 国家\u3000\u3000\u3000", 11, False, BODY),
        ("q\u2082", 11, True, RED), (":8 月 \u00d7 加拿大 \u00d7 产品\u3000\u3000\u3000", 11, False, BODY),
        ("q\u2083", 11, True, RED), (":8 月 \u00d7 加拿大 \u00d7 客户类型 \u2014\u2014 三条查询同一数据子集,仅聚合维度不同", 11, False, BODY),
    ]])

    tb(s, 0.67, 1.80, 5.8, 0.24, [("方案 A · 逐条独立优化(现状)", 12, True, INK)])
    tb(s, 6.87, 1.80, 5.8, 0.24, [("方案 B · 轨迹感知优化(本课题)", 12, True, RED)])

    tA = make_table(s, 0.67, 2.08, 5.8, [0.55, 3.85, 1.4], [0.3] * 5)
    rowsA = [
        ("查询", "最优计划", "代价", None),
        ("q\u2081", "全表扫描 + 过滤 + 聚合", "100", BODY),
        ("q\u2082", "全表扫描 + 过滤 + 聚合", "100", BODY),
        ("q\u2083", "全表扫描 + 过滤 + 聚合", "100", BODY),
        ("合计", "同一子集被过滤了三遍", "300", INK),
    ]
    for i, (a, b, c, cc) in enumerate(rowsA):
        head = i == 0
        fill = None if head else (PALE if i == 4 else (PALE if i % 2 == 0 else WHITE))
        set_cell(tA, i, 0, [(a, True, RED if head else INK)], 10 if head else 9.5, fill, PP_ALIGN.CENTER)
        set_cell(tA, i, 1, [(b, head, RED if head else BODY)], 9.5 if not head else 10, fill)
        set_cell(tA, i, 2, [(c, True, RED if head else (INK if i == 4 else BODY))], 10 if head else 11, fill, PP_ALIGN.CENTER)

    tB = make_table(s, 6.87, 2.08, 5.8, [0.55, 3.85, 1.4], [0.3] * 5)
    rowsB = [
        ("查询", "计划", "代价", None),
        ("q\u2081", "扫描 + 顺带物化 M\u2081(2 GB)", "110", BODY),
        ("q\u2082", "直接扫 M\u2081(q\u2082 \u2286 M\u2081)", "12", BODY),
        ("q\u2083", "直接扫 M\u2081(换聚合维度)", "12", BODY),
        ("合计", "整体成本 \u2193 55%", "134", RED),
    ]
    for i, (a, b, c, cc) in enumerate(rowsB):
        head = i == 0
        fill = None if head else (PALE if i % 2 == 0 else WHITE)
        set_cell(tB, i, 0, [(a, True, RED if head else INK)], 10 if head else 9.5, fill, PP_ALIGN.CENTER)
        set_cell(tB, i, 1, [(b, head, RED if head else BODY)], 9.5 if not head else 10, fill)
        set_cell(tB, i, 2, [(c, True, RED if head else cc)], 10 if head else 11, fill, PP_ALIGN.CENTER)

    tb(s, 0.67, 3.68, 5.8, 0.2, [("三条各自最优、无一可再省 \u2014\u2014 但整体在重复劳动", 9.5, False, MUTE)])
    tb(s, 6.87, 3.68, 5.8, 0.2, [("物化让 q\u2081 变贵 10(110 > 100),换回 q\u2082、q\u2083 各省 88", 9.5, True, RED)])

    args = [
        ("\u2460 ", "C(q\u2081 + 物化) = 110 > C(q\u2081) = 100 \u2014\u2014 物化对 q\u2081 单独是亏本决策,单查询优化器永远不会主动做;"),
        ("\u2461 ", "决策发生在执行 q\u2081 的时刻:q\u2082、q\u2083 尚未产生(Agent 还没看到 r\u2081)\u2014\u2014 没有未来输入,也没有 batch 可看;"),
        ("\u2462 ", "8 月子集在 q\u2081 执行时明明流过内存,却即用即弃 \u2014\u2014 缺的不是数据,是跨查询物理状态机制。"),
    ]
    yy = 3.96
    for lead, txt in args:
        tb(s, 0.67, yy, 12.0, 0.22, [[(lead, 10, True, RED), (txt, 10, False, BODY)]], space=0)
        yy += 0.26

    navy_bar(s, 4.86, [[
        ("\u03a3 min Cost(q\u1d62) = 300    \u2260    min Cost(q\u2081, q\u2082, q\u2083) = 134", 12.5, True, WHITE),
        ("    \u2192 目标升级为 min\u209c Cost(\u03c4, \u03c0)", 12.5, True, RGBColor(0xFF, 0xB3, 0xB3)),
    ]], h=0.44)

    sec_head(s, 0.67, 5.5, 5.8, "Challenge|为什么难")
    facts = [
        ("下一条 SQL 此刻不存在", " \u2014\u2014 决策时 q\u209c\u208a\u2081 尚未产生,优化器没有未来输入"),
        ("未来由当前结果决定", " \u2014\u2014 q\u209c\u208a\u2081 = Agent(r\u209c, State\u209c),r\u209c 未知则方向未知"),
        ("意图 \u2260 SQL", " \u2014\u2014 意图不保证兑现,也不唯一对应具体查询"),
        ("单查询目标错位", " \u2014\u2014 跨查询收益不在任何单条查询的目标函数内"),
        ("准备必须先花钱", " \u2014\u2014 收益是概率性的,代价是即时且确定的"),
        ("状态不能无限堆", " \u2014\u2014 每条查询都留中间结果会压垮系统"),
    ]
    yy = 5.86
    for lead, txt in facts:
        tb(s, 0.67, yy, 5.8, 0.2, [[(lead, 10, True, INK), (txt, 10, False, BODY)]], space=0)
        yy += 0.22

    sec_head(s, 6.87, 5.5, 5.8, "Opportunity|为什么有机会", red=True)
    opps = [
        ("Goal-Driven", " \u2014\u2014 多条查询围绕同一高层任务,方向可预期"),
        ("Intent-Exposable", " \u2014\u2014 SQL 未生成,Agent 意图已可提前暴露"),
        ("子集 / 聚合相关", " \u2014\u2014 q\u2082 \u2286 q\u2081、q\u2083 = Agg(q\u2082) 等关系大量存在"),
        ("中间结果在流动", " \u2014\u2014 子集流过内存,只是即用即弃"),
        ("动作空间统一", " \u2014\u2014 物化 / 索引 / 预取 / 保留 / 淘汰,皆为物理状态"),
        ("上游生态就绪", " \u2014\u2014 Cortex Agents / Genie 等可供给机器可读意图"),
    ]
    yy = 5.86
    for lead, txt in opps:
        tb(s, 6.87, yy, 5.8, 0.2, [[(lead, 10, True, RED), (txt, 10, False, BODY)]], space=0)
        yy += 0.22
    footer(s, "4")


def page4():
    s = prs.slides.add_slide(BLANK)
    header(s, "业务价值:从 Query Latency 到 Agent Task Completion Time", "03 · 业务价值",
           "学术新问题 \u00d7 工程真问题 \u2014\u2014 既有新的信息条件与优化范式,也有真实的成本、规模与体验压力")

    tb(s, 0.67, 1.46, 12.0, 0.55, [
        ("Agent 正成为数据库新的、高增长的数据访问主体:一个任务产生数条至数十条相关 SQL;云端按量计费(credits / Request Unit)下,重复计算就是直接浪费。", 11.5, False, BODY),
    ], space=0)

    sections = [
        (0.67, 2.14, "\u2460 端到端任务提速", [
            [("优化目标跃迁:", 10.5, True, INK), ("Query Latency \u2192 Trajectory Cost \u2192 Task Completion Time", 10.5, True, RED)],
            ("消除轨迹中的重复扫描 / Join / 聚合,多轮分析整体提速;Agent 迭代越快,洞察交付越快。", 10, False, BODY),
        ]),
        (6.84, 2.14, "\u2461 计算与 Token 成本双降", [
            ("查询更快 \u2192 Agent 每轮等待更短 \u2192 相同预算完成更多分析轮次。", 10.5, False, BODY),
            ("参照:腾讯云 AI 诊断将 CPU 异常定位从 30 分钟压缩到 2 分钟。", 10, False, MUTE),
        ]),
        (0.67, 4.28, "\u2462 Agent 规模化的基础设施", [
            ("上线规模 = 并发 Agent 数 \u00d7 每 SQL 成本,内核吞吐是天花板。", 10.5, False, BODY),
            ("\u201c数据库的主要用户不再是人类,而是 AI Agent\u201d \u2014\u2014 PingCAP", 10, False, MUTE),
        ]),
        (6.84, 4.28, "\u2463 与 Agent 产品生态互补", [
            ("上游解决「生成正确的 SQL」,本课题解决「高效执行整条轨迹」。", 10.5, False, BODY),
            ("物化 / 索引 / 预取决策在内核完成,对 Agent 透明,可嵌入任何 Data Agent 产品。", 10, False, BODY),
        ]),
    ]
    for x, y, title, paras in sections:
        sec_head(s, x, y, 5.83, title)
        tb(s, x, y + 0.42, 5.83, 1.4, paras, space=4)

    tb(s, 0.67, 6.32, 12.0, 0.26, [[
        ("关键判断:", 12, True, RED),
        ("这是与模型能力提升正交、可以由系统侧直接获取的性能收益 \u2014\u2014 不依赖更强的模型,只依赖更懂 Agent 的数据库。", 12, True, INK),
    ]])
    navy_bar(s, 6.72, [[
        ("落地场景:ChatBI / Deep Research 归因 · 企业级 Data Agent · 自动化运维 Agent · 多租户 SaaS 分析", 11.5, True, WHITE),
    ]])
    footer(s, "5")


def page5():
    s = prs.slides.add_slide(BLANK)
    header(s, "已有系统如何获得「未来工作负载」信息:四种范式", "04 · 相关工作",
           "按「数据库做当前决策时,对未来 workload 知道什么」归类")

    t = make_table(s, 0.67, 1.42, 12.0, [1.75, 2.45, 2.45, 2.45, 2.9], [0.42] + [0.78] * 4)
    heads = ["维度", "\u2160 Current-Query\n无未来信息", "\u2161 Known-Future\n未来查询已知", "\u2162 Inferred-Future\n从历史推断未来", "\u2163 Exposed-Future\n意图可暴露 \u2605 本课题"]
    for c, h in enumerate(heads):
        set_cell(t, 0, c, [(h, True, RED if c == 4 else RED)], 10.5, None, PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER, line=1.1)

    body_rows = [
        ("信息条件", "当前查询 + Schema + 统计信息", "q\u209c\u208a\u2081, q\u209c\u208a\u2082, \u2026 直接可见(显式 + 确定)", "P(W_f | W_past) \u2014 隐式 + 概率性", "q\u209c\u208a\u2081 未知,但 Goal / Plan / Stage / CandidateActions 可暴露"),
        ("回答的问题", "\u201c这条查询怎么执行最便宜?\u201d", "\u201c给定整批 workload,如何共享执行?\u201d", "\u201c根据过去的行为,接下来可能访问什么?\u201d", "\u201c负载生成者接下来想做什么,我现在该准备什么?\u201d"),
        ("代表系统", "传统 Query Optimizer", "MQO \u00b7 ReProVide \u00b7 AutoCox(procedural SQL)", "HAWC \u00b7 HashStash \u00b7 SparkCruise \u00b7 ForeCache \u00b7 PreView", "本课题:Agent\u2013DB 协同优化"),
        ("特点", "Reactive:查询到达 \u2192 优化 \u2192 执行", "Perfect Foresight:直接规划跨查询共享", "Behavior-Driven:由观察反推未来", "Query Unknown, Intent Visible"),
    ]
    for r, row in enumerate(body_rows):
        for c, val in enumerate(row):
            if c == 0:
                fill = None
                set_cell(t, r + 1, c, [(val, True, INK)], 10.5, fill)
            else:
                ours = c == 4
                fill = PALE if ours else (PALE if r % 2 == 1 else WHITE)
                set_cell(t, r + 1, c, [(val, ours, RED if ours else BODY)], 9.5, fill, line=1.15)

    tb(s, 0.67, 5.28, 12.0, 0.26, [[
        ("关键差异不在「是否预测未来」,而在「未来信息从哪里来」:", 12, True, INK),
        ("观察行为推断 \u2192 生成者主动暴露", 12, True, RED),
    ]], align=PP_ALIGN.CENTER)
    tb(s, 0.67, 5.62, 12.0, 0.24, [
        ("范式演进:Reactive(只看当前查询)\u2192 Predictive(从历史预测未来)\u2192 Cooperative(意图 + 轨迹观测协同)", 11, False, BODY),
    ], align=PP_ALIGN.CENTER)
    tb(s, 0.67, 6.28, 12.0, 0.2, [
        ("注:PreView(2026)、AutoCox(2026, Data & Knowledge Engineering)为近年工作,正式版将补充完整引用。", 8.5, False, LIGHT),
    ])
    navy_bar(s, 6.72, [[
        ("信息条件升级:P(W_f | W_past)  \u2192  P(W_f | W_past, R_past, AgentContext)", 12, True, WHITE),
        ("    \u2014\u2014 未来意图第一次由 workload generator 主动提供", 11.5, True, RGBColor(0xFF, 0xB3, 0xB3)),
    ]])
    footer(s, "6")


def page5b():
    s = prs.slides.add_slide(BLANK)
    header(s, "研究现状全景:与六类相关工作的系统对比", "04 · 相关工作",
           "16 个比较维度 \u00d7 6 个研究方向;\u2713 具备 / \u2717 不具备 / \u25b3 部分具备或视情况,红色列为本研究")

    REDTINT = RGBColor(0xFA, 0xEC, 0xEC)
    heads = ["比较维度", "Traditional Query Optimization", "MQO / Known Query Sequence", "History-Aware Reuse",
             "Interactive / Predictive Exploration", "Agentic Query / Workflow Systems",
             "本研究：Agent-Aware Trajectory Optimization"]

    rows_data = [
        ("主要优化单元", [("单条 Query", False), ("Query batch / 已知 sequence", False), ("Query history + current query", False), ("Interactive session / future access", False), ("Agentic plan / workflow", False), ("Evolving Agent query trajectory", True)], 0.31),
        ("未来具体 Query 是否可见", [("\u2717", False), ("\u2713", True), ("\u2717", False), ("\u2717", False), ("\u25b3，取决于 workflow", False), ("\u2717", True)], 0.285),
        ("未来访问是否可预测", [("\u2717", False), ("不需要预测", False), ("\u25b3，从历史 reuse pattern 推断", False), ("\u2713，从历史/交互行为预测", True), ("\u25b3", False), ("\u2713", True)], 0.285),
        ("Future intent 是否直接可见", [("\u2717", False), ("Query 本身已知，因此通常不需要 intent", False), ("\u2717", False), ("\u2717，用户 intent 通常是黑盒", True), ("\u25b3 / \u2713", False), ("\u2713，作为核心输入", True)], 0.31),
        ("Workload generator 是否参与优化", [("\u2717", False), ("\u2717", False), ("\u2717", False), ("通常 \u2717", False), ("\u2713，Agent 属于执行系统", False), ("\u2713，Agent 与 DB 显式协同", True)], 0.285),
        ("Future query 是否受当前结果影响", [("不考虑", False), ("通常已固定", False), ("\u25b3", False), ("\u2713", True), ("\u2713", True), ("\u2713，核心 workload 特征", True)], 0.285),
        ("Future uncertainty 是否显式建模", [("\u2717", False), ("\u2717", False), ("\u25b3", False), ("\u2713", False), ("\u25b3", False), ("\u2713，作为优化问题的一部分", True)], 0.285),
        ("跨 Query 共享计算/结果", [("\u2717", False), ("\u2713", True), ("\u2713", True), ("mechanism-specific", False), ("通常不是核心", False), ("\u2713", True)], 0.285),
        ("主动创建 Future-Oriented Physical State", [("\u2717", False), ("\u2713", False), ("\u25b3 / \u2713", False), ("\u2713，通常针对特定机制", False), ("\u25b3", False), ("\u2713，核心能力", True)], 0.285),
        ("物化结果 / Intermediate Reuse", [("当前 Query 内", False), ("\u2713", False), ("\u2713", True), ("部分", False), ("\u25b3", False), ("\u2713", True)], 0.285),
        ("Prefetch", [("execution-level", False), ("部分", False), ("部分", False), ("\u2713", True), ("部分", False), ("\u2713", True)], 0.285),
        ("Ephemeral Physical Design / Index", [("现有 index selection", False), ("部分，如 procedural SQL", False), ("较少", False), ("较少", False), ("非核心", False), ("\u2713，潜在统一 action", True)], 0.285),
        ("多类 Physical State 的统一决策", [("\u2717", False), ("\u25b3", False), ("\u25b3", False), ("通常单机制", False), ("\u2717", False), ("目标：\u2713", True)], 0.285),
        ("核心 Future 信息来源", [("无", False), ("未来 Query 本身", True), ("Past Queries", True), ("Past Queries / User Behavior", True), ("Agent workflow state", False), ("Agent Intent + Trajectory Observation", True)], 0.31),
        ("信息状态", [("No Future", False), ("Known Future", True), ("Inferred Future", True), ("Inferred Future", True), ("Agent-Aware", False), ("Exposed but Uncertain Future", True)], 0.285),
        ("核心优化目标", [("Current-query cost", False), ("Batch / sequence cost", False), ("Future query reuse benefit", False), ("Interactive latency", False), ("Quality / workflow cost", False), ("Expected trajectory-level DB cost", True)], 0.31),
    ]

    heights = [0.48] + [r[2] for r in rows_data]
    t = make_table(s, 0.67, 1.40, 12.0, [1.85] + [1.69] * 6, heights)
    for c, h in enumerate(heads):
        set_cell(t, 0, c, [(h, True, RED)], 8.5, None, PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER, line=1.0)
    for i, (label, cells, _) in enumerate(rows_data):
        fill = WHITE if i % 2 == 0 else PALE
        set_cell(t, i + 1, 0, [(label, True, INK)], 7.5, fill)
        for c, (txt, bold) in enumerate(cells):
            if c == 5:
                color, cf = RED, REDTINT
            elif txt == "\u2717":
                color, cf = LIGHT, fill
            elif txt.startswith("\u25b3"):
                color, cf = MUTE, fill
            else:
                color, cf = BODY, fill
            set_cell(t, i + 1, c + 1, [(txt, bold, color)], 8, cf, line=1.05)

    navy_bar(s, 6.72, [[
        ("最关键的差异不是「本研究一列 \u2713 最多」,而是前三个未来信息维度 \u2014\u2014 ", 11.5, True, WHITE),
        ("Future Query Visibility / Predictability / Intent Visibility", 11.5, True, RGBColor(0xFF, 0xB3, 0xB3)),
        (" \u2014\u2014 它们定义了根本的信息条件差异", 11.5, True, WHITE),
    ]])
    footer(s, "7")


def page6():
    s = prs.slides.add_slide(BLANK)
    header(s, "研究空白:Future Query Unknown, Future Intent Visible", "04 · 研究定位",
           "两个维度定位全部相关工作:未来查询可见性 \u00d7 负载生成者可见性")

    rect(s, 1.05, 1.5, 2.6, 2.12, PALE, RED, 1.75)
    quads = [
        (1.2, 1.62, [("\u2605 本课题", 13, True, RED), ("Query Unknown, Intent Visible", 10.5, True, INK),
                     ("Agent 暴露意图 + 轨迹观测", 9.5, False, BODY), ("\u2192 跨查询物理状态优化", 9.5, False, BODY)]),
        (3.85, 1.62, [("计划完全确定的 Agent 工作流", 10.5, True, MUTE), ("query 已知 + 生成者可见(退化 / 少见情形)", 9.5, False, LIGHT)]),
        (1.2, 3.82, [("Inferred-Future", 11, True, INK), ("EDA \u00b7 ForeCache \u00b7 HAWC \u00b7 SparkCruise \u00b7 PreView", 9.5, False, BODY), ("观察行为 \u2192 推断未来(黑盒生成者)", 9.5, False, MUTE)]),
        (3.85, 3.82, [("Known-Future", 11, True, INK), ("MQO \u00b7 ReProVide \u00b7 AutoCox", 9.5, False, BODY), ("未来 query 直接可见(Perfect Foresight)", 9.5, False, MUTE)]),
    ]
    for x, y, paras in quads:
        tb(s, x, y, 2.5, 1.0, paras, space=2, line=1.15)

    rect(s, 1.0, 5.95, 5.55, 0.018, INK)
    tri = rect(s, 6.5, 5.865, 0.16, 0.18, INK, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
    tri.rotation = 90
    tb(s, 1.0, 6.06, 5.55, 0.22, [("Future Query Visibility:Unknown \u2192 Known", 10, True, INK)], align=PP_ALIGN.CENTER)

    rect(s, 1.0, 1.45, 0.018, 4.5, INK)
    tri2 = rect(s, 0.925, 1.42, 0.16, 0.18, INK, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
    ybox = tb(s, 0.28, 2.6, 0.42, 2.2, [("Workload Generator Visibility:Black-box \u2192 Intent Visible", 10, True, INK)],
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    ybox.text_frame._txBody.bodyPr.set("vert", "vert270")

    tb(s, 0.67, 6.42, 6.2, 0.4, [
        ("Agentic systems 让 workload generator 第一次变得可观察,但其优化对象是 Agent 的动作流,而非 DB 的跨查询物理状态。", 9.5, False, MUTE),
    ], space=0)

    tb(s, 7.35, 1.42, 5.3, 0.26, [("与最近邻的三个本质差异", 13, True, INK)])
    hrule(s, 7.35, 1.72, 5.3)
    diffs = [
        ("\u2460 Intent before Query", [
            ("未来 SQL 生成之前,Agent 已可暴露 Goal / Stage / Candidate Actions;信息来源从「观察行为」变为「生成者主动供给」。", 10.5, False, BODY)]),
        ("\u2461 Uncertain by Nature", [
            ("Intent \u2260 Future SQL:后续查询仍由中间结果驱动,可能随时转向;不确定性必须显式进入优化问题。", 10.5, False, BODY)]),
        ("\u2462 Optimize Physical State", [
            ("数据库将意图转化为 retain / materialize / prefetch / index / evict 决策 \u2014\u2014", 10.5, False, BODY),
            ("Agent-Aware Cross-Query Physical State Management", 10.5, True, RED)]),
    ]
    yy = 1.9
    for title, paras in diffs:
        tb(s, 7.35, yy, 5.3, 0.26, [(title, 12.5, True, RED)])
        tb(s, 7.35, yy + 0.3, 5.3, 0.85, paras, space=3, line=1.2)
        yy += 1.42

    navy_bar(s, 6.72, [[
        ("Known \u2192 Inferred \u2192 Exposed Future      |      Reactive \u2192 Predictive \u2192 Cooperative      |      ", 11.5, True, WHITE),
        ("Query Unknown, Intent Visible", 12, True, RGBColor(0xFF, 0xB3, 0xB3)),
    ]])
    footer(s, "8")


cover()
page1()
page2()
page3()
page4()
page5()
page5b()
page6()

out = "Agent-Aware-Trajectory-Level-Query-Processing.pptx"
prs.save(out)
print("saved:", out, "slides:", len(prs.slides._sldIdLst))
